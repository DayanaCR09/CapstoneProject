import json
import re
from collections.abc import Mapping
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from exceptions import DataValidationError

_INVALID_SHEET_CHARACTERS = re.compile(r"[\[\]:*?/\\]")
_INVALID_FILENAME_CHARACTERS = re.compile(r"[<>:\"/\\|?*\x00-\x1f]")


class ReportGenerator:
    """Export pandas datasets as formatted Excel, CSV, and JSON reports."""

    _HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
    _HEADER_FONT = Font(color="FFFFFF", bold=True)
    _HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")

    @staticmethod
    def _validate_dataframe(dataframe: pd.DataFrame, name: str = "dataframe") -> None:
        if not isinstance(dataframe, pd.DataFrame):
            raise DataValidationError(f"{name} must be a pandas DataFrame")

    @classmethod
    def _validate_datasets(
        cls,
        datasets: Mapping[str, pd.DataFrame],
    ) -> None:
        if not isinstance(datasets, Mapping) or not datasets:
            raise DataValidationError(
                "datasets must be a non-empty mapping of names to DataFrames",
                field="datasets",
            )

        sheet_names: set[str] = set()
        for name, dataframe in datasets.items():
            if not isinstance(name, str) or not name.strip():
                raise DataValidationError(
                    "dataset names must be non-empty strings",
                    field="datasets",
                )
            normalized_sheet_name = cls._sheet_name(name)
            if normalized_sheet_name.casefold() in sheet_names:
                raise DataValidationError(
                    "dataset names must be unique after Excel sheet-name normalization",
                    field="datasets",
                )
            sheet_names.add(normalized_sheet_name.casefold())
            cls._validate_dataframe(dataframe, f"dataset {name!r}")

    @staticmethod
    def _sheet_name(name: str) -> str:
        sheet_name = _INVALID_SHEET_CHARACTERS.sub("_", name.strip())[:31]
        if not sheet_name:
            raise DataValidationError(
                "dataset name cannot produce a valid Excel sheet name",
                field="datasets",
            )
        return sheet_name

    @staticmethod
    def _filename_component(name: str) -> str:
        component = _INVALID_FILENAME_CHARACTERS.sub("_", name.strip()).strip(" .")
        if not component:
            raise DataValidationError(
                "dataset name cannot produce a valid filename",
                field="datasets",
            )
        return component

    @staticmethod
    def _excel_value(value: Any) -> Any:
        if value is None:
            return None
        if not pd.api.types.is_scalar(value):
            return str(value)
        try:
            if bool(pd.isna(value)):
                return None
        except TypeError, ValueError:
            pass

        if isinstance(value, pd.Timestamp):
            return value.to_pydatetime()
        if isinstance(value, (datetime, date, str, int, float, bool)):
            return value
        if hasattr(value, "item"):
            value = value.item()
            if isinstance(value, (datetime, date, str, int, float, bool)):
                return value
        return str(value)

    @classmethod
    def _write_sheet(
        cls,
        worksheet: Worksheet,
        dataframe: pd.DataFrame,
    ) -> None:
        worksheet.append([str(column) for column in dataframe.columns])
        for cell in worksheet[1]:
            cell.fill = cls._HEADER_FILL
            cell.font = cls._HEADER_FONT
            cell.alignment = cls._HEADER_ALIGNMENT

        for row in dataframe.itertuples(index=False, name=None):
            worksheet.append([cls._excel_value(value) for value in row])
            for cell in worksheet[worksheet.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"

        worksheet.freeze_panes = "A2"
        if dataframe.columns.size:
            worksheet.auto_filter.ref = worksheet.dimensions
        for column_index, column in enumerate(dataframe.columns, start=1):
            header = str(column)
            values = (
                worksheet.cell(row=row_index, column=column_index).value
                for row_index in range(2, worksheet.max_row + 1)
            )
            content_width = max(
                (len(str(value)) for value in values if value is not None),
                default=0,
            )
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(len(header), content_width) + 2,
                60,
            )

    @staticmethod
    def _prepare_path(output_path: str | Path) -> Path:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        return path

    @classmethod
    def export_excel(
        cls,
        datasets: Mapping[str, pd.DataFrame],
        output_path: str | Path,
    ) -> Path:
        """Write one named DataFrame per formatted worksheet in an Excel workbook."""
        cls._validate_datasets(datasets)
        path = cls._prepare_path(output_path)

        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, dataframe in datasets.items():
            worksheet = workbook.create_sheet(cls._sheet_name(name))
            cls._write_sheet(worksheet, dataframe)
        workbook.save(path)
        return path

    @classmethod
    def export_csv_summary(
        cls,
        dataframe: pd.DataFrame,
        output_path: str | Path,
    ) -> Path:
        """Write a dataset as a structured CSV table without its pandas index."""
        cls._validate_dataframe(dataframe)
        path = cls._prepare_path(output_path)
        dataframe.to_csv(path, index=False, encoding="utf-8")
        return path

    @classmethod
    def export_json(
        cls,
        dataframe: pd.DataFrame,
        output_path: str | Path,
        *,
        minified: bool = False,
    ) -> Path:
        """Write a dataset as a UTF-8 JSON array of records."""
        cls._validate_dataframe(dataframe)
        path = cls._prepare_path(output_path)
        records = json.loads(
            dataframe.to_json(
                orient="records",
                date_format="iso",
                force_ascii=False,
            )
        )
        path.write_text(
            json.dumps(
                records,
                ensure_ascii=False,
                indent=None if minified else 2,
                separators=(",", ":") if minified else None,
                allow_nan=False,
            )
            + "\n",
            encoding="utf-8",
        )
        return path

    @classmethod
    def export_reports(
        cls,
        datasets: Mapping[str, pd.DataFrame],
        output_directory: str | Path,
        *,
        filename_prefix: str = "report",
    ) -> dict[str, Path]:
        """Write a combined workbook and CSV/JSON files for every dataset."""
        cls._validate_datasets(datasets)
        if not isinstance(filename_prefix, str) or not filename_prefix.strip():
            raise DataValidationError(
                "filename_prefix must be a non-empty string",
                field="filename_prefix",
            )

        prefix = cls._filename_component(filename_prefix)
        filenames: set[str] = set()
        components = {}
        for name in datasets:
            component = cls._filename_component(name)
            filename = f"{prefix}_{component}".casefold()
            if filename in filenames:
                raise DataValidationError(
                    "dataset names must be unique after filename normalization",
                    field="datasets",
                )
            filenames.add(filename)
            components[name] = component

        output_directory = Path(output_directory)
        output_directory.mkdir(parents=True, exist_ok=True)
        outputs = {
            "excel": cls.export_excel(
                datasets,
                output_directory / f"{prefix}.xlsx",
            )
        }
        for name, dataframe in datasets.items():
            component = components[name]
            outputs[f"{name}_csv"] = cls.export_csv_summary(
                dataframe,
                output_directory / f"{prefix}_{component}.csv",
            )
            outputs[f"{name}_json"] = cls.export_json(
                dataframe,
                output_directory / f"{prefix}_{component}.json",
            )
        return outputs
