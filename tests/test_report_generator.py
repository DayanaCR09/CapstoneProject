import csv
import json
from unittest.mock import call, patch

import pandas as pd
import pytest
from openpyxl import load_workbook

from exceptions import DataValidationError
from report_generator import ReportGenerator


def test_exports_formatted_workbook_with_sheets_and_safe_cell_values(tmp_path):
    dataframe = pd.DataFrame(
        {
            "title": ["Pride and Prejudice", "=1+1", None],
            "published": pd.to_datetime(["1813-01-28", None, "2024-01-01"]),
            "copies": [3, 2, 1],
        }
    )
    path = tmp_path / "nested" / "catalog.xlsx"

    result = ReportGenerator.export_excel(
        {"Cleaned works": dataframe, "By author": dataframe[["title", "copies"]]},
        path,
    )

    assert result == path
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["Cleaned works", "By author"]
    worksheet = workbook["Cleaned works"]
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:C4"
    assert worksheet["A1"].font.bold
    assert worksheet["A1"].fill.fgColor.rgb.endswith("1F4E78")
    assert worksheet["A3"].value == "=1+1"
    assert worksheet["A3"].data_type == "s"
    assert worksheet["A4"].value is None
    assert worksheet["B2"].is_date
    assert worksheet.column_dimensions["A"].width > len("title")


def test_exports_csv_without_dataframe_index(tmp_path):
    dataframe = pd.DataFrame(
        {"title": ["Pride and Prejudice", "Emma"], "count": [3, 1]},
        index=[10, 20],
    )
    path = ReportGenerator.export_csv_summary(dataframe, tmp_path / "summary.csv")

    with path.open(encoding="utf-8", newline="") as csv_file:
        rows = list(csv.reader(csv_file))

    assert rows == [["title", "count"], ["Pride and Prejudice", "3"], ["Emma", "1"]]


def test_exports_json_records_and_nulls(tmp_path):
    dataframe = pd.DataFrame(
        {
            "title": ["Crème brûlée", None],
            "published": pd.to_datetime(["2024-01-02", None]),
            "count": [2, 0],
        }
    )
    path = ReportGenerator.export_json(dataframe, tmp_path / "summary.json")

    assert json.loads(path.read_text(encoding="utf-8")) == [
        {
            "title": "Crème brûlée",
            "published": "2024-01-02T00:00:00.000",
            "count": 2,
        },
        {"title": None, "published": None, "count": 0},
    ]


def test_export_reports_writes_workbook_and_dataset_summaries(tmp_path):
    datasets = {
        "Cleaned works": pd.DataFrame({"title": ["Emma"]}),
        "By author": pd.DataFrame({"author": ["Jane Austen"], "count": [1]}),
    }

    outputs = ReportGenerator.export_reports(
        datasets,
        tmp_path,
        filename_prefix="books",
    )

    assert set(outputs) == {
        "excel",
        "Cleaned works_csv",
        "Cleaned works_json",
        "By author_csv",
        "By author_json",
    }
    assert all(path.exists() for path in outputs.values())


def test_export_reports_delegates_to_export_apis_with_expected_paths(tmp_path):
    datasets = {
        "Cleaned works": pd.DataFrame({"title": ["Emma"]}),
        "By author": pd.DataFrame({"author": ["Jane Austen"], "count": [1]}),
    }
    excel_path = tmp_path / "books.xlsx"
    cleaned_csv = tmp_path / "books_Cleaned works.csv"
    cleaned_json = tmp_path / "books_Cleaned works.json"
    author_csv = tmp_path / "books_By author.csv"
    author_json = tmp_path / "books_By author.json"

    with (
        patch.object(ReportGenerator, "export_excel", return_value=excel_path) as excel,
        patch.object(
            ReportGenerator,
            "export_csv_summary",
            side_effect=[cleaned_csv, author_csv],
        ) as csv_export,
        patch.object(
            ReportGenerator,
            "export_json",
            side_effect=[cleaned_json, author_json],
        ) as json_export,
    ):
        outputs = ReportGenerator.export_reports(
            datasets,
            tmp_path,
            filename_prefix="books",
        )

    assert outputs == {
        "excel": excel_path,
        "Cleaned works_csv": cleaned_csv,
        "Cleaned works_json": cleaned_json,
        "By author_csv": author_csv,
        "By author_json": author_json,
    }
    excel.assert_called_once_with(datasets, tmp_path / "books.xlsx")
    assert csv_export.call_args_list == [
        call(datasets["Cleaned works"], tmp_path / "books_Cleaned works.csv"),
        call(datasets["By author"], tmp_path / "books_By author.csv"),
    ]
    assert json_export.call_args_list == [
        call(datasets["Cleaned works"], tmp_path / "books_Cleaned works.json"),
        call(datasets["By author"], tmp_path / "books_By author.json"),
    ]


@pytest.mark.parametrize(
    "datasets",
    [
        {},
        {"": pd.DataFrame()},
        {"invalid/name": pd.DataFrame(), "invalid?name": pd.DataFrame()},
    ],
)
def test_export_excel_rejects_invalid_or_duplicate_sheet_names(datasets, tmp_path):
    with pytest.raises(DataValidationError):
        ReportGenerator.export_excel(datasets, tmp_path / "invalid.xlsx")


def test_export_rejects_non_dataframe_input(tmp_path):
    with pytest.raises(DataValidationError, match="DataFrame"):
        ReportGenerator.export_csv_summary(
            [{"title": "Emma"}],
            tmp_path / "invalid.csv",
        )
